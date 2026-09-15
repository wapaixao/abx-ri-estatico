#!/usr/bin/env python3
"""Constrói a DFC indireta preliminar ABX 1T26/2T26 a partir do data.json.

Premissas:
- Resultado parte da DRE formal.
- BP 001-011 permanece no bloco Água Branca matriz/filiais.
- Disponibilidade é a política conservadora de caixa; aplicações ficam no FCI.
- Variações de BP são líquidas; não se inventam movimentos brutos.
- Diferenças de conciliação ficam explícitas e não integram FCO/FCI/FCF.
"""
from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
DATA_PATH = BASE / "data.json"
AUDIT_DIR = BASE / "audit_hidden"
DELIVERY_DIR = Path("/root/data/abx/entregas/APRESENTACAO")
AUDIT_XLSX = AUDIT_DIR / "ABX_DFC_1T_2T_2026_PRELIMINAR_AUDITORIA.xlsx"
DELIVERY_XLSX = DELIVERY_DIR / "ABX_DFC_1T_2T_2026_PRELIMINAR.xlsx"
TOLERANCE = 1.0

PERIODS: Dict[str, Tuple[str, str]] = {
    "1T26": ("31/12/2025", "31/03/2026"),
    "2T26": ("31/03/2026", "30/06/2026"),
}

ENTITIES = [
    ("Água Branca matriz/filiais", "TOTAL FILIAIS 001-011"),
    ("HortiVan", "050 - Hortivan"),
    ("Top Morena / Top Frutas", "100 - Top Frutas"),
    ("Água Branca Maringá", "101 - Água Branca"),
    ("Top Verde", "103 - Top Verde"),
]

# (índice no BP, rótulo no DFC)
OPERATING_ASSETS = [
    (3, "Redução / (aumento) de estoques"),
    (4, "Redução / (aumento) de créditos a receber"),
    (5, "Redução / (aumento) de recebíveis antecipados"),
    (6, "Redução / (aumento) de adiantamentos e outros ativos"),
    (10, "Redução / (aumento) de renegociações de longo prazo"),
    (11, "Redução / (aumento) de créditos de ICMS"),
    (12, "Redução / (aumento) de créditos de PIS e COFINS"),
]
OPERATING_LIABILITIES = [
    (18, "Aumento / (redução) de fornecedores"),
    (19, "Aumento / (redução) de aluguéis a pagar"),
    (20, "Aumento / (redução) de salários e contribuições"),
    (21, "Aumento / (redução) de impostos a pagar"),
    (24, "Aumento / (redução) de adiantamentos passivos"),
    (25, "Variação de reclassificações do passivo circulante"),
    (30, "Aumento / (redução) de débitos tributários de longo prazo"),
]
INVESTING_ASSETS = [
    (8, "Aplicações e seguros fora de equivalentes — variação líquida"),
    (9, "Capitalizações — variação líquida"),
    (13, "Imobilizado — variação líquida"),
    (14, "Obras em bens de terceiros — variação líquida"),
    (15, "Intangível — variação líquida"),
]
FINANCING_LIABILITIES = [
    (22, "Empréstimos e financiamentos de curto prazo — variação líquida"),
    (23, "Créditos rotativos — variação líquida"),
    (27, "Banco do Brasil de longo prazo — variação líquida"),
    (28, "Bradesco de longo prazo — variação líquida"),
    (29, "Arrendamentos / leasing de longo prazo — variação líquida"),
    (31, "Outras obrigações não circulantes — variação líquida"),
]


def row_values(companies: Iterable[str], periods: Iterable[str]) -> dict:
    return {company: {period: 0.0 for period in periods} for company in companies}


def main() -> None:
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    bp = data["reports"]["BP"]
    dre = data["reports"]["DRE"]
    if bp.get("pl_audit_status") != "audited":
        raise RuntimeError("PL não está marcado como auditado; DFC bloqueada.")

    bp_rows = bp["rows"]
    bp_by_label = {r["descricao"]: r for r in bp_rows}
    dre_profit = next(r for r in dre["rows"] if r["descricao"] == "LUCRO LIQUIDO")
    companies = [bp_name for bp_name, _ in ENTITIES]

    definitions = []
    definitions.append(("FCO", 1, "Fluxo de caixa das atividades operacionais (FCO)", "subtotal"))
    definitions.append(("LL", 3, "Lucro líquido — DRE", "profit"))
    definitions.extend((f"OA_{idx}", 3, label, "asset_decrease") for idx, label in OPERATING_ASSETS)
    definitions.extend((f"OL_{idx}", 3, label, "liability_increase") for idx, label in OPERATING_LIABILITIES)
    definitions.append(("FCI", 1, "Fluxo de caixa das atividades de investimento (FCI)", "subtotal"))
    definitions.extend((f"IA_{idx}", 3, label, "asset_decrease") for idx, label in INVESTING_ASSETS)
    definitions.append(("FCF", 1, "Fluxo de caixa das atividades de financiamento (FCF)", "subtotal"))
    definitions.extend((f"FL_{idx}", 3, label, "liability_increase") for idx, label in FINANCING_LIABILITIES)
    definitions.append(("PL_EX_LL", 3, "Movimentos líquidos de PL, exceto lucro da DRE — a detalhar", "equity_ex_profit"))
    definitions.append(("CASH_CHANGE", 1, "Variação líquida de caixa", "cash_change"))
    definitions.append(("CASH_START", 2, "Caixa e equivalentes no início", "cash_start"))
    definitions.append(("CASH_END", 2, "Caixa e equivalentes no final", "cash_end"))
    definitions.append(("RECON", 2, "Diferença de conciliação", "reconciliation"))

    values = {key: row_values(companies, PERIODS) for key, _, _, _ in definitions}
    status = row_values(companies, PERIODS)

    op_asset_keys = [f"OA_{i}" for i, _ in OPERATING_ASSETS]
    op_liab_keys = [f"OL_{i}" for i, _ in OPERATING_LIABILITIES]
    inv_keys = [f"IA_{i}" for i, _ in INVESTING_ASSETS]
    fin_keys = [f"FL_{i}" for i, _ in FINANCING_LIABILITIES]

    for bp_name, dre_name in ENTITIES:
        for period, (opening, closing) in PERIODS.items():
            if dre.get("coverage", {}).get(dre_name, {}).get(period) is False:
                raise RuntimeError(f"DRE sem fonte: {dre_name} {period}")
            profit = float(dre_profit["empresas"][dre_name][period])
            values["LL"][bp_name][period] = profit
            for idx, _ in OPERATING_ASSETS:
                values[f"OA_{idx}"][bp_name][period] = float(
                    bp_rows[idx]["empresas"][bp_name][opening] - bp_rows[idx]["empresas"][bp_name][closing]
                )
            for idx, _ in OPERATING_LIABILITIES:
                values[f"OL_{idx}"][bp_name][period] = float(
                    bp_rows[idx]["empresas"][bp_name][closing] - bp_rows[idx]["empresas"][bp_name][opening]
                )
            values["FCO"][bp_name][period] = profit + sum(
                values[key][bp_name][period] for key in op_asset_keys + op_liab_keys
            )

            for idx, _ in INVESTING_ASSETS:
                values[f"IA_{idx}"][bp_name][period] = float(
                    bp_rows[idx]["empresas"][bp_name][opening] - bp_rows[idx]["empresas"][bp_name][closing]
                )
            values["FCI"][bp_name][period] = sum(values[key][bp_name][period] for key in inv_keys)

            for idx, _ in FINANCING_LIABILITIES:
                values[f"FL_{idx}"][bp_name][period] = float(
                    bp_rows[idx]["empresas"][bp_name][closing] - bp_rows[idx]["empresas"][bp_name][opening]
                )
            pl_change = float(
                bp_by_label["PATRIMÔNIO LÍQUIDO"]["empresas"][bp_name][closing]
                - bp_by_label["PATRIMÔNIO LÍQUIDO"]["empresas"][bp_name][opening]
            )
            values["PL_EX_LL"][bp_name][period] = pl_change - profit
            values["FCF"][bp_name][period] = sum(values[key][bp_name][period] for key in fin_keys) + values["PL_EX_LL"][bp_name][period]

            start = float(bp_by_label["Disponibilidade"]["empresas"][bp_name][opening])
            end = float(bp_by_label["Disponibilidade"]["empresas"][bp_name][closing])
            change = values["FCO"][bp_name][period] + values["FCI"][bp_name][period] + values["FCF"][bp_name][period]
            reconciliation = end - (start + change)
            values["CASH_START"][bp_name][period] = start
            values["CASH_CHANGE"][bp_name][period] = change
            values["CASH_END"][bp_name][period] = end
            values["RECON"][bp_name][period] = reconciliation
            status[bp_name][period] = "CONCILIADO" if abs(reconciliation) <= TOLERANCE else "PENDENTE"

    rows = []
    for key, level, label, nature in definitions:
        row = {
            "codigo": key,
            "secao": "DFC",
            "nivel": level,
            "descricao": label,
            "natureza": nature,
            "empresas": values[key],
            "grupo": {p: sum(values[key][c][p] for c in companies) for p in PERIODS},
        }
        rows.append(row)

    group_status = {
        p: "CONCILIADO" if abs(next(r for r in rows if r["codigo"] == "RECON")["grupo"][p]) <= TOLERANCE else "PENDENTE"
        for p in PERIODS
    }
    dfc = {
        "label": "DFC — Demonstração dos Fluxos de Caixa (preliminar)",
        "periods": list(PERIODS),
        "companies": bp["companies"],
        "rows": rows,
        "method": "Método indireto, partindo do lucro líquido da DRE e das variações do BP contábil",
        "status": status,
        "group_status": group_status,
        "cash_policy": "Disponibilidade do BP; aplicações, seguros, capitalizações e SWAP não classificados como equivalentes sem evidência CPC 03.",
        "granularity": "Água Branca matriz/filiais (001-011 consolidado), 050, 100, 101-M, 103 e Grupo.",
        "limitations": [
            "Movimentos patrimoniais apresentados por variação líquida; sem razão, não há segregação de fluxos brutos.",
            "Movimentos de PL, exceto o lucro da DRE, permanecem agrupados até separar aportes, distribuições, pagamentos e reclassificações.",
            "O ajuste gerencial de R$ 108.550 exibido no relatório específico de PL não integra esta DFC; a DFC usa exclusivamente o BP contábil e a DRE formal.",
            "Transferências internas entre 001-011 não são apresentadas como fluxo e ficam absorvidas no bloco consolidado de tesouraria.",
        ],
        "coverage_note": "1T26 e 2T26; unidades 004 e 010 mantidas como sem movimento confirmado no 2T26.",
        "sources": {
            "BP": list(bp.get("sources", [])),
            "DRE": list(dre.get("sources", [])),
        },
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "published": False,
    }
    data["reports"]["DFC"] = dfc

    backup_dir = AUDIT_DIR / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"data.backup_pre_dfc_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    shutil.copy2(DATA_PATH, backup)
    DATA_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    build_workbook(dfc)
    DELIVERY_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(AUDIT_XLSX, DELIVERY_XLSX)

    recon = next(r for r in rows if r["codigo"] == "RECON")
    print(f"DFC gravada em {DATA_PATH}")
    print(f"Backup: {backup.name}")
    print(f"Planilha: {DELIVERY_XLSX}")
    for company in companies:
        print(company, {p: recon["empresas"][company][p] for p in PERIODS}, status[company])
    print("Grupo", recon["grupo"], group_status)


def build_workbook(dfc: dict) -> None:
    AUDIT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "DFC Preliminar"
    dark = "123D20"
    green = "D8E9A8"
    yellow = "FFF2CC"
    red = "F4CCCC"
    white = "FFFFFF"
    line = Side(style="thin", color="C8D6B5")

    companies = [c["name"] for c in dfc["companies"]]
    ws.append(["DFC ABX — método indireto preliminar"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + len(companies) * len(PERIODS))
    ws["A1"].fill = PatternFill("solid", fgColor=dark)
    ws["A1"].font = Font(color=white, bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="left")

    header1 = ["Descrição"]
    header2 = [""]
    for company in companies:
        header1.extend([company, ""])
        header2.extend(PERIODS.keys())
    ws.append(header1)
    ws.append(header2)
    col = 2
    for _ in companies:
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)
        col += 2
    for row in ws.iter_rows(min_row=2, max_row=3):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=dark)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(left=line, right=line, top=line, bottom=line)

    row_map = {}
    for source_row in dfc["rows"]:
        excel_row = ws.max_row + 1
        row_map[source_row["codigo"]] = excel_row
        vals = [source_row["descricao"]]
        for company in companies:
            vals.extend(source_row["empresas"][company][p] for p in PERIODS)
        ws.append(vals)
        fill = green if source_row["nivel"] == 1 else white
        if source_row["codigo"] == "RECON":
            fill = red
        for cell in ws[excel_row]:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(bold=source_row["nivel"] == 1 or source_row["codigo"] == "RECON")
            cell.border = Border(left=line, right=line, top=line, bottom=line)
        for cell in ws[excel_row][1:]:
            cell.number_format = '#,##0;[Red]-#,##0;–'
            cell.alignment = Alignment(horizontal="right")

    ws.freeze_panes = "B4"
    ws.column_dimensions["A"].width = 66
    for c in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16

    rec = wb.create_sheet("Reconciliação")
    rec.append(["Bloco", "Período", "Caixa inicial", "FCO", "FCI", "FCF", "Diferença", "Caixa final", "Status"])
    by_code = {r["codigo"]: r for r in dfc["rows"]}
    for company in companies:
        for period in PERIODS:
            rec.append([
                company, period,
                by_code["CASH_START"]["empresas"][company][period],
                by_code["FCO"]["empresas"][company][period],
                by_code["FCI"]["empresas"][company][period],
                by_code["FCF"]["empresas"][company][period],
                by_code["RECON"]["empresas"][company][period],
                by_code["CASH_END"]["empresas"][company][period],
                dfc["status"][company][period],
            ])
    rec.append([
        "TOTAL GRUPO", "2T26",
        by_code["CASH_START"]["grupo"]["2T26"], by_code["FCO"]["grupo"]["2T26"],
        by_code["FCI"]["grupo"]["2T26"], by_code["FCF"]["grupo"]["2T26"],
        by_code["RECON"]["grupo"]["2T26"], by_code["CASH_END"]["grupo"]["2T26"],
        dfc["group_status"]["2T26"],
    ])
    for cell in rec[1]:
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.font = Font(color=white, bold=True)
    for row in rec.iter_rows(min_row=2):
        for cell in row[2:8]:
            cell.number_format = '#,##0;[Red]-#,##0;–'
        if row[8].value == "PENDENTE":
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=red)
    rec.freeze_panes = "A2"
    for col, width in {"A": 34, "B": 12, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16, "I": 14}.items():
        rec.column_dimensions[col].width = width

    notes = wb.create_sheet("Metodologia e Fontes")
    note_rows = [
        ("Método", dfc["method"]),
        ("Política de caixa", dfc["cash_policy"]),
        ("Granularidade", dfc["granularity"]),
        ("Cobertura", dfc["coverage_note"]),
        ("Publicação", "Não publicada; prévia local para revisão."),
    ]
    note_rows.extend((f"Limitação {i}", note) for i, note in enumerate(dfc["limitations"], 1))
    def source_text(source: object) -> str:
        if isinstance(source, dict):
            return " | ".join(f"{key}: {value}" for key, value in source.items())
        return str(source)

    note_rows.extend((f"Fonte BP {i}", source_text(src)) for i, src in enumerate(dfc["sources"]["BP"], 1))
    note_rows.extend((f"Fonte DRE {i}", source_text(src)) for i, src in enumerate(dfc["sources"]["DRE"], 1))
    notes.append(["Item", "Descrição"])
    for item, desc in note_rows:
        notes.append([item, desc])
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.font = Font(color=white, bold=True)
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 120
    for row in notes.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(AUDIT_XLSX)


if __name__ == "__main__":
    main()
